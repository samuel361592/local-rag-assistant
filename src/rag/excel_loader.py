"""Read XLSX values and cached formulas as bounded, complete-row documents."""

from __future__ import annotations

from collections import Counter
from contextlib import ExitStack
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from math import isfinite
from pathlib import Path
from typing import BinaryIO, Iterable
from zipfile import is_zipfile

from langchain_core.documents import Document
from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from openpyxl.styles.numbers import is_datetime
from openpyxl.utils import get_column_letter

from .config import Settings
from .processing_errors import ExcelProcessingError


@dataclass(frozen=True, slots=True)
class ExcelExtractionStats:
    workbook_count: int = 1
    total_sheet_count: int = 0
    indexed_sheet_count: int = 0
    skipped_empty_sheet_count: int = 0
    indexed_row_count: int = 0
    formula_without_cache_count: int = 0


@dataclass(frozen=True, slots=True)
class ExcelRow:
    row_number: int
    text: str


def _format_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip().replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        if not isfinite(value):
            raise ExcelProcessingError("儲存格含非有限數值。")
        return "0" if value == 0 else format(value, ".15g")
    return str(value)


def _check_cell_length(text: str, coordinate: str, settings: Settings) -> None:
    if len(text) > settings.excel_max_cell_chars:
        raise ExcelProcessingError(
            f"儲存格 {coordinate} 超過 EXCEL_MAX_CELL_CHARS={settings.excel_max_cell_chars}。"
        )


def _cell_value_text(cell: ReadOnlyCell | EmptyCell) -> str:
    value = cell.value
    # Excel's date-only serials are loaded as midnight datetime objects.
    if isinstance(value, datetime) and is_datetime(getattr(cell, "number_format", "")) == "date":
        value = value.date()
    return _format_value(value)


def format_excel_cell(
    value_cell: ReadOnlyCell | EmptyCell, formula_cell: ReadOnlyCell | EmptyCell,
    settings: Settings,
) -> tuple[str, bool]:
    no_cache = formula_cell.data_type == "f" and value_cell.value is None
    coordinate = getattr(formula_cell, "coordinate", "未知")
    text = (
        f"無可用計算值（公式儲存格 {coordinate}）"
        if no_cache else _cell_value_text(value_cell)
    )
    _check_cell_length(text, coordinate, settings)
    return text, no_cache


def _unique_name(candidate: str, column: int, used: set[str]) -> str:
    if candidate in used:
        candidate += f" [欄 {get_column_letter(column)}]"
    base = candidate
    number = 2
    while candidate in used:
        candidate = f"{base} #{number}"
        number += 1
    used.add(candidate)
    return candidate


def pack_excel_rows(
    rows: Iterable[ExcelRow], sheet_name: str, base_metadata: dict, settings: Settings,
) -> list[Document]:
    prefix = f"工作表：{sheet_name}\n\n"
    target = min(settings.chunk_size, settings.excel_max_chunk_chars)
    chunks: list[Document] = []
    current: list[ExcelRow] = []
    current_size = len(prefix)

    def flush() -> None:
        nonlocal current_size
        if not current:
            return
        content = prefix + "\n".join(row.text for row in current)
        if len(content) > settings.excel_max_chunk_chars:
            raise ExcelProcessingError("Excel Chunk 超過 EXCEL_MAX_CHUNK_CHARS。")
        chunks.append(Document(page_content=content, metadata={
            **base_metadata,
            "row_start": current[0].row_number,
            "row_end": current[-1].row_number,
            "oversized_row": len(current) == 1 and len(content) > settings.chunk_size,
        }))
        current.clear()
        current_size = len(prefix)

    for row in rows:
        if len(row.text) > settings.excel_max_row_chars:
            raise ExcelProcessingError(
                f"第 {row.row_number} 列超過 EXCEL_MAX_ROW_CHARS={settings.excel_max_row_chars}。"
            )
        if len(prefix) + len(row.text) > settings.excel_max_chunk_chars:
            raise ExcelProcessingError(
                f"第 {row.row_number} 列超過 EXCEL_MAX_CHUNK_CHARS={settings.excel_max_chunk_chars}。"
            )
        if current and current_size + 1 + len(row.text) > target:
            flush()
        current_size += len(row.text) + bool(current)
        current.append(row)
        if current_size > target:
            flush()
    flush()
    return chunks


def load_excel_documents(
    uploaded_file: BinaryIO, settings: Settings,
) -> tuple[list[Document], ExcelExtractionStats]:
    filename = Path(getattr(uploaded_file, "name", "未命名.xlsx")).name
    sheet_name = ""
    try:
        if Path(filename).suffix.casefold() != ".xlsx":
            raise ExcelProcessingError("僅支援 XLSX；不支援 XLS 或 XLSM。")
        if hasattr(uploaded_file, "getvalue"):
            content = uploaded_file.getvalue()
        else:
            uploaded_file.seek(0)
            content = uploaded_file.read(settings.excel_max_file_size_mb * 1024 * 1024 + 1)
        if len(content) > settings.excel_max_file_size_mb * 1024 * 1024:
            raise ExcelProcessingError(
                f"檔案超過 EXCEL_MAX_FILE_SIZE_MB={settings.excel_max_file_size_mb} MB。"
            )
        if content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            raise ExcelProcessingError("不支援需要開啟密碼的加密活頁簿或舊版 Excel 格式。")
        if not is_zipfile(BytesIO(content)):
            raise ExcelProcessingError("檔案可能損毀或不是有效 XLSX／ZIP。")
        with ExitStack() as stack:
            values = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
            stack.callback(values.close)
            formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            stack.callback(formulas.close)
            sheets = values.worksheets
            if len(sheets) > settings.excel_max_sheets:
                raise ExcelProcessingError(f"工作表超過 EXCEL_MAX_SHEETS={settings.excel_max_sheets}。")
            chunks: list[Document] = []
            indexed_sheets = row_count = missing_cache_count = 0
            for sheet_index, value_sheet in enumerate(sheets):
                sheet_name = value_sheet.title
                if len(sheet_name) > 31:
                    raise ExcelProcessingError("工作表名稱超過 31 字元。")
                formula_sheet = formulas[sheet_name]
                for sheet in (value_sheet, formula_sheet):
                    if (sheet.max_row or 0) > settings.excel_max_row_index:
                        raise ExcelProcessingError(f"列座標超過 EXCEL_MAX_ROW_INDEX={settings.excel_max_row_index}。")
                    if (sheet.max_column or 0) > settings.excel_max_column_index:
                        raise ExcelProcessingError(f"欄座標超過 EXCEL_MAX_COLUMN_INDEX={settings.excel_max_column_index}。")
                    # Do not lose cells when a producer writes understated dimensions.
                    sheet.reset_dimensions()
                column_names: dict[int, str] = {}
                used_names: set[str] = set()
                metadata = {
                    "source": filename, "source_filename": filename,
                    "source_type": "excel", "content_type": "text",
                    "text_extraction_method": "excel", "sheet_name": sheet_name,
                    "sheet_index": sheet_index,
                    "sheet_hidden": value_sheet.sheet_state != "visible",
                }

                def rows() -> Iterable[ExcelRow]:
                    nonlocal row_count, missing_cache_count
                    value_rows, formula_rows = value_sheet.iter_rows(), formula_sheet.iter_rows()
                    try:
                        for row_number, (vr, fr) in enumerate(zip(value_rows, formula_rows, strict=True), 1):
                            if row_number > settings.excel_max_row_index:
                                raise ExcelProcessingError(f"列座標超過 EXCEL_MAX_ROW_INDEX={settings.excel_max_row_index}。")
                            if max(len(vr), len(fr)) > settings.excel_max_column_index:
                                raise ExcelProcessingError(f"欄座標超過 EXCEL_MAX_COLUMN_INDEX={settings.excel_max_column_index}。")
                            pairs = list(zip(vr, fr, strict=True))
                            missing_cache_count += sum(f.data_type == "f" and v.value is None for v, f in pairs)
                            if "header_row" not in metadata:
                                names = [" ".join(_cell_value_text(v).split()) for v, _ in pairs]
                                if not any(names):
                                    continue
                                metadata["header_row"] = row_number
                                counts = Counter(name for name in names if name)
                                for col, name in enumerate(names, 1):
                                    letter = get_column_letter(col)
                                    candidate = (f"{name} [{letter}]" if counts[name] > 1 else name) if name else f"未命名欄位 [{letter}]"
                                    column_names[col] = _unique_name(candidate, col, used_names)
                                    _check_cell_length(column_names[col], f"{letter}{row_number}", settings)
                                continue
                            fields = []
                            for col, (v, f) in enumerate(pairs, 1):
                                text, _ = format_excel_cell(v, f, settings)
                                if not text:
                                    continue
                                if col not in column_names:
                                    column_names[col] = _unique_name(f"未命名欄位 [{get_column_letter(col)}]", col, used_names)
                                    _check_cell_length(column_names[col], f"{get_column_letter(col)}{metadata['header_row']}", settings)
                                fields.append(f"{column_names[col]}：{text}")
                            if fields:
                                row_count += 1
                                yield ExcelRow(row_number, f"第 {row_number} 列｜" + "｜".join(fields))
                    finally:
                        value_rows.close()
                        formula_rows.close()

                records = rows()
                try:
                    sheet_chunks = pack_excel_rows(records, sheet_name, metadata, settings)
                finally:
                    records.close()
                indexed_sheets += bool(sheet_chunks)
                chunks.extend(sheet_chunks)
            if not chunks:
                raise ExcelProcessingError("沒有可建立索引的資料列。")
            return chunks, ExcelExtractionStats(
                total_sheet_count=len(sheets), indexed_sheet_count=indexed_sheets,
                skipped_empty_sheet_count=len(sheets) - indexed_sheets,
                indexed_row_count=row_count, formula_without_cache_count=missing_cache_count,
            )
    except ExcelProcessingError as exc:
        location = f"，工作表「{sheet_name}」" if sheet_name else ""
        raise ExcelProcessingError(f"Excel「{filename}」{location}：{exc}") from None
    except Exception:
        raise ExcelProcessingError(f"Excel「{filename}」可能損毀或格式不正確，無法解析。") from None
