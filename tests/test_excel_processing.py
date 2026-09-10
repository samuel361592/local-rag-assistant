from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile
import re

import pytest
from langchain_core.documents import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from src.rag import document_loader, excel_loader
from src.rag.config import Settings, load_settings
from src.rag.document_loader import DocumentExtractionStats, process_uploaded_documents
from src.rag.excel_loader import ExcelRow, format_excel_cell, load_excel_documents, pack_excel_rows
from src.rag.processing_errors import DocumentProcessingError, ExcelProcessingError


class Upload(BytesIO):
    def __init__(self, content: bytes, name: str = "sales.xlsx"):
        super().__init__(content)
        self.name = name


def upload(workbook: Workbook, name="sales.xlsx") -> Upload:
    stream = BytesIO()
    workbook.save(stream)
    return Upload(stream.getvalue(), name)


def basic_workbook() -> Workbook:
    wb = Workbook()
    wb.active.title = "Sales"
    wb.active.append(["產品", "金額"])
    wb.active.append(["A", 20])
    return wb


def rewrite_xml(file: Upload, transform) -> Upload:
    """Synthetic XML fixtures, deliberately not attributed to Excel/LibreOffice."""
    output = BytesIO()
    with ZipFile(file) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename.startswith("xl/worksheets/sheet"):
                data = transform(data.decode()).encode()
            target.writestr(entry, data)
    return Upload(output.getvalue(), file.name)


def test_workbook_headers_hidden_empty_rows_and_metadata():
    wb = Workbook()
    ws = wb.active
    ws.title = "資料"
    ws.append([" ", "=1+1"])
    ws.append(["名稱", "名稱", "名稱 [A]", None])
    ws.append(["甲", "乙", "丙", "丁"])
    ws.append([None])
    ws.append(["戊", None, None, None, "表頭外"])
    hidden = wb.create_sheet("隱藏")
    hidden.sheet_state = "veryHidden"
    hidden.append(["欄名"])
    hidden.append(["=1+1"])
    blank = wb.create_sheet("空白")
    blank["A1"] = "=1+1"
    chunks, stats = load_excel_documents(upload(wb), Settings())
    assert stats.total_sheet_count == 3
    assert stats.indexed_sheet_count == 2
    assert stats.skipped_empty_sheet_count == 1
    assert stats.indexed_row_count == 3
    assert stats.formula_without_cache_count == 3
    first = chunks[0]
    assert first.metadata == {
        "source": "sales.xlsx", "source_filename": "sales.xlsx", "source_type": "excel",
        "content_type": "text", "text_extraction_method": "excel", "sheet_name": "資料",
        "sheet_index": 0, "sheet_hidden": False, "header_row": 2, "row_start": 3, "row_end": 5,
        "oversized_row": False,
    }
    assert "名稱 [A]：甲｜名稱 [B]：乙｜名稱 [A] [欄 C]：丙｜未命名欄位 [D]：丁" in first.page_content
    assert "第 5 列｜名稱 [A]：戊｜未命名欄位 [E]：表頭外" in first.page_content
    assert "第 4 列" not in first.page_content
    assert chunks[-1].metadata["sheet_hidden"] is True
    assert "無可用計算值（公式儲存格 A2）" in chunks[-1].page_content


def test_dynamic_names_remain_fixed_with_collision():
    wb = Workbook()
    ws = wb.active
    ws.append(["未命名欄位 [E]", " 名稱\n  X ", "名稱 X"])
    ws.cell(2, 5, "one")
    ws.cell(3, 5, "two")
    chunks, _ = load_excel_documents(upload(wb), Settings())
    assert "未命名欄位 [E] [欄 E]：one" in chunks[0].page_content
    assert "未命名欄位 [E] [欄 E]：two" in chunks[0].page_content


@pytest.mark.parametrize("value, expected", [
    (None, ""), ("  A\r\nB\nC  ", "A B C"), (True, "TRUE"), (False, "FALSE"),
    (12345678901234567890, "12345678901234567890"), (1.2300, "1.23"),
    (0.000001, "1e-06"), (1e20, "1e+20"), (-0.0, "0"),
    (date(2026, 9, 8), "2026-09-08"), (datetime(2026, 9, 8, 1, 2, 3), "2026-09-08 01:02:03"),
    (time(1, 2, 3), "01:02:03"), ("#N/A", "#N/A"),
])
def test_cell_formats(value, expected):
    cell = SimpleNamespace(value=value, data_type="s", coordinate="A2")
    assert format_excel_cell(cell, cell, Settings()) == (expected, False)


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -float("inf")])
def test_nonfinite_rejected(value):
    cell = SimpleNamespace(value=value, data_type="n", coordinate="A2")
    with pytest.raises(ExcelProcessingError):
        format_excel_cell(cell, cell, Settings())


def test_date_and_datetime_roundtrip_preserves_distinct_formats():
    wb = Workbook()
    wb.active.append(["date", "datetime", "time"])
    wb.active.append([date(2026, 9, 8), datetime(2026, 9, 8, 0, 0, 0), time(12, 30)])
    chunks, _ = load_excel_documents(upload(wb), Settings())
    assert "date：2026-09-08｜datetime：2026-09-08 00:00:00｜time：12:30:00" in chunks[0].page_content


def test_cached_formula_and_cached_header_synthetic():
    wb = Workbook()
    wb.active.append(['="標題"', "數值"])
    wb.active.append(["A", "=10*2"])
    file = rewrite_xml(upload(wb), lambda xml: re.sub(
        r'<c r="A1">.*?</c>',
        '<c r="A1" t="str"><f>"標題"</f><v>標題</v></c>',
        xml,
    ).replace('<f>10*2</f><v></v>', '<f>10*2</f><v>20</v>'))
    chunks, stats = load_excel_documents(file, Settings())
    assert "標題：A｜數值：20" in chunks[0].page_content
    assert "10*2" not in chunks[0].page_content
    assert stats.formula_without_cache_count == 0


def test_formula_original_is_not_subject_to_value_limit():
    wb = basic_workbook()
    wb.active["B2"] = "=" + "1+" * 100 + "1"
    _, stats = load_excel_documents(upload(wb), Settings(excel_max_cell_chars=30))
    assert stats.formula_without_cache_count == 1


@pytest.mark.parametrize("row, col, styled", [(50001, 1, False), (50001, 1, True), (2, 201, False), (2, 201, True)])
def test_coordinate_limits(row, col, styled):
    wb = basic_workbook()
    cell = wb.active.cell(row, col)
    if styled:
        cell.font = Font(bold=True)
    else:
        cell.value = "value"
    with pytest.raises(ExcelProcessingError, match="EXCEL_MAX_.*_INDEX") as error:
        load_excel_documents(upload(wb), Settings())
    assert "sales.xlsx" in str(error.value) and "Sales" in str(error.value)


def test_boundary_row_and_understated_dimensions():
    wb = basic_workbook()
    wb.active["A50000"] = "last"
    file = rewrite_xml(upload(wb), lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:A1"', xml))
    chunks, stats = load_excel_documents(file, Settings())
    assert stats.indexed_row_count == 2
    assert chunks[-1].metadata["row_end"] == 50000


def test_understated_dimensions_do_not_bypass_limits():
    wb = basic_workbook()
    wb.active["A50001"] = "last"
    file = rewrite_xml(upload(wb), lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:A1"', xml))
    with pytest.raises(ExcelProcessingError, match="EXCEL_MAX_ROW_INDEX"):
        load_excel_documents(file, Settings())


@pytest.mark.parametrize("kind", ["empty", "header", "formula_only"])
def test_no_data_workbook_rejected(kind):
    wb = Workbook()
    if kind == "header":
        wb.active["A1"] = "header"
    if kind == "formula_only":
        wb.active["A1"] = "=1+1"
    with pytest.raises(ExcelProcessingError, match="沒有可建立索引"):
        load_excel_documents(upload(wb), Settings())


def test_merged_values_not_filled_and_protection_accepted():
    wb = basic_workbook()
    ws = wb.active
    ws.merge_cells("A2:A3")
    ws["B3"] = 30
    ws.protection.sheet = True
    wb.security.lockStructure = True
    chunks, _ = load_excel_documents(upload(wb), Settings())
    assert "第 3 列｜金額：30" in chunks[0].page_content


@pytest.mark.parametrize("kind", ["cell", "header", "row", "sheets", "size"])
def test_resource_limits_reject_with_safe_messages(kind):
    wb = basic_workbook()
    settings = Settings()
    if kind == "cell":
        wb.active["A2"] = "SECRET" * 600
    elif kind == "header":
        wb.active["A1"] = "SECRET" * 600
    elif kind == "row":
        wb.active["A2"] = "x" * 2500
        wb.active["B2"] = "y" * 2500
    elif kind == "sheets":
        wb.create_sheet("extra")
        settings = replace(settings, excel_max_sheets=1)
    else:
        file = Upload(b"x" * (1024 * 1024 + 1))
        settings = replace(settings, excel_max_file_size_mb=1)
    with pytest.raises(ExcelProcessingError) as error:
        load_excel_documents(file if kind == "size" else upload(wb), settings)
    assert "sales.xlsx" in str(error.value)
    assert "SECRET" not in str(error.value)


@pytest.mark.parametrize("content", [b"garbage", bytes.fromhex("D0CF11E0A1B11AE1")])
def test_invalid_or_encrypted_file_safe(content):
    with pytest.raises(ExcelProcessingError) as error:
        load_excel_documents(Upload(content), Settings())
    assert "Traceback" not in str(error.value)


def test_workbook_closed_when_second_load_fails(monkeypatch):
    workbook = SimpleNamespace(close=lambda: closed.append(True))
    closed = []
    calls = []
    def fake_load(*args, **kwargs):
        calls.append(kwargs)
        if len(calls) == 2:
            raise ValueError("secret internals")
        return workbook
    monkeypatch.setattr(excel_loader, "load_workbook", fake_load)
    with pytest.raises(ExcelProcessingError) as error:
        load_excel_documents(upload(basic_workbook()), Settings())
    assert closed == [True]
    assert "secret" not in str(error.value)


def test_packing_limits_and_oversized():
    settings = Settings(chunk_size=30, chunk_overlap=0, excel_max_row_chars=80, excel_max_chunk_chars=117)
    rows = [ExcelRow(2, "a" * 10), ExcelRow(5, "b" * 10), ExcelRow(6, "c" * 40)]
    chunks = pack_excel_rows(rows, "S", {}, settings)
    assert [c.metadata["row_start"] for c in chunks] == [2, 6]
    assert chunks[0].metadata["row_end"] == 5
    assert chunks[1].metadata["oversized_row"] is True
    assert chunks[1].page_content.endswith("c" * 40)
    larger = replace(settings, chunk_size=500)
    chunks = pack_excel_rows([ExcelRow(i, "x" * 40) for i in range(5)], "S", {}, larger)
    assert all(len(c.page_content) <= 117 for c in chunks)
    assert all(not c.metadata["oversized_row"] for c in chunks)
    with pytest.raises(ExcelProcessingError, match="ROW_CHARS"):
        pack_excel_rows([ExcelRow(2, "x" * 81)], "S", {}, settings)
    with pytest.raises(ExcelProcessingError, match="CHUNK_CHARS"):
        pack_excel_rows([ExcelRow(2, "x" * 80)], "s" * 40, {}, settings)


def test_excel_only_does_not_initialize_models_or_splitter(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Excel must not call OCR, VLM or the text splitter")
    for name in ("OCRService", "VLMService", "split_pdf_pages"):
        monkeypatch.setattr(document_loader, name, forbidden)
    second = basic_workbook()
    second.active["A2"] = "different source"
    batch = process_uploaded_documents([
        upload(basic_workbook(), "a.xlsx"), upload(second, "b.xlsx"),
    ], Settings(vlm_enabled=True))
    assert batch.page_count == 0
    assert batch.document_count == 2
    assert batch.extraction_stats.excel_workbook_count == 2
    assert batch.extraction_stats.excel_total_sheet_count == 2
    assert batch.extraction_stats.excel_indexed_row_count == 2


@pytest.mark.parametrize("others", [[], ["manual.pdf"], ["image.png"], ["manual.pdf", "image.png"]])
def test_mixed_batch_combines_all_sources(monkeypatch, others):
    def fake_load(file, settings, ocr, vlm):
        source_type = "pdf" if file.name.endswith("pdf") else "image"
        return ([Document(page_content="sample", metadata={"source_type": source_type, "text_extraction_method": "ocr" if source_type == "image" else "native"})], [], DocumentExtractionStats(), {}, {}, {})
    monkeypatch.setattr(document_loader, "_load_pdf_pages", fake_load)
    monkeypatch.setattr(document_loader, "_load_image_page", fake_load)
    batch = process_uploaded_documents([upload(basic_workbook()), *[Upload(name.encode(), name) for name in others]], Settings())
    assert len(batch.chunks) == 1 + len(others)
    assert batch.page_count == len(others)
    assert any(c.metadata["source_type"] == "excel" for c in batch.chunks)


def test_duplicate_and_unsupported_rejected_before_parsing():
    with pytest.raises(DocumentProcessingError, match="重複"):
        process_uploaded_documents([Upload(b"", "sales.xlsx"), Upload(b"", "Sales.xlsx")], Settings())
    with pytest.raises(DocumentProcessingError, match="PDF、PNG、JPG、JPEG、XLSX"):
        process_uploaded_documents([Upload(b"", "data.xls")], Settings())


def test_excel_settings_env_and_validation(tmp_path, monkeypatch):
    monkeypatch.setenv("EXCEL_MAX_ROW_CHARS", "4000")
    monkeypatch.setenv("EXCEL_MAX_CHUNK_CHARS", "4096")
    env = tmp_path / ".env"
    env.write_text("EXCEL_MAX_ROW_CHARS=100\nEXCEL_MAX_CHUNK_CHARS=137\n", encoding="utf-8")
    settings = load_settings(env_path=env, override_env=True)
    assert settings.excel_max_row_chars == 100
    assert settings.excel_max_chunk_chars == 137
    for name in Settings.__dataclass_fields__:
        if name.startswith("excel_"):
            with pytest.raises(ValueError):
                replace(Settings(), **{name: 0})
    with pytest.raises(ValueError, match="37"):
        Settings(excel_max_chunk_chars=4036)


FIXTURES = Path(__file__).parent / "fixtures" / "excel"


@pytest.mark.parametrize("filename", ["formula_cached_synthetic.xlsx", "formula_cached.xlsx"])
def test_fixed_formula_cache_fixture(filename):
    path = FIXTURES / filename
    if filename == "formula_cached.xlsx" and not path.exists():
        pytest.skip("Requires maintainer-supplied Excel/LibreOffice-saved formula_cached.xlsx; see fixtures README")
    values = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    try:
        assert values["FormulaCache"]["D2"].value == 20
        assert values["FormulaCache"]["D3"].value == 21
        assert formulas["FormulaCache"]["D2"].value == "=B2*C2"
        assert formulas["FormulaCache"]["D3"].value == "=B3*C3"
    finally:
        values.close()
        formulas.close()
    chunks, stats = load_excel_documents(Upload(path.read_bytes()), Settings())
    assert stats.formula_without_cache_count == 0
    assert "Amount：20" in chunks[0].page_content


def test_success_and_failure_fixture_files():
    chunks, stats = load_excel_documents(Upload((FIXTURES / "excel_success_cases.xlsx").read_bytes()), Settings())
    assert stats.total_sheet_count == 3
    assert stats.indexed_sheet_count == 2
    assert stats.formula_without_cache_count == 2
    assert any(chunk.metadata["oversized_row"] for chunk in chunks)
    for path in FIXTURES.glob("failure_*.xlsx"):
        with pytest.raises(ExcelProcessingError):
            load_excel_documents(Upload(path.read_bytes(), path.name), Settings())


def test_real_mixed_loaders_and_faiss_with_fake_model_dependencies():
    import fitz
    from PIL import Image
    from langchain_core.embeddings import Embeddings
    from src.rag.ocr_service import OCRResult
    from src.rag.vlm_service import VisualAnalysisResult
    from src.rag.vector_store import create_vector_store

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Maintenance manual: check equipment before operation every day.")
    pdf_content = pdf.tobytes()
    pdf.close()
    png = BytesIO()
    Image.new("RGB", (30, 20), "white").save(png, format="PNG")
    ocr_calls, vlm_calls = [], []
    class OCR:
        def image_bytes_to_text(self, data, *, suffix):
            ocr_calls.append(suffix)
            return OCRResult("equipment MX-20", 0.9)
    class VLM:
        def analyze_image(self, data, *, mime_type, filename, page_number):
            vlm_calls.append(filename)
            return VisualAnalysisResult(page_type="文件", summary="可見設備檢查說明。")
    class LocalEmbeddings(Embeddings):
        def embed_documents(self, texts):
            self.texts = texts
            return [[float(i), 1.0] for i, _ in enumerate(texts)]
        def embed_query(self, text):
            return [0.0, 1.0]
    settings = Settings(vlm_enabled=True, vlm_mode="all", ocr_dpi=72)
    batch = process_uploaded_documents([
        Upload(pdf_content, "manual.pdf"), Upload(png.getvalue(), "image.png"),
        upload(basic_workbook()),
    ], settings, ocr_service=OCR(), vlm_service=VLM())
    assert ocr_calls == [".png"]
    assert vlm_calls == ["manual.pdf", "image.png"]
    assert batch.page_count == 2
    assert batch.extraction_stats.visual_chunk_count == 2
    embedding = LocalEmbeddings()
    store = create_vector_store(batch.chunks, settings, embeddings=embedding)
    assert embedding.texts == [chunk.page_content for chunk in batch.chunks]
    assert store.index.ntotal == len(batch.chunks)
    retrieved = store.similarity_search("sample", k=len(batch.chunks))
    assert {doc.metadata["source_type"] for doc in retrieved} == {"pdf", "image", "excel"}
